"""
Competency editing is allowed only before the syllabus enters the approval chain;
the VP's final endorsement validates the reviewed list and locks it.
"""
import uuid

import pytest

from backend.core.security import create_access_token, hash_password
from backend.models.course import Course
from backend.models.syllabus import ApprovalStage, Competency, Syllabus, SyllabusStatus
from backend.models.user import User, UserRole


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _approver(db, role, email):
    user = User(email=email, password_hash=hash_password("Approver@1234"), full_name=email,
                role=role, is_active=True, college="ICS")
    db.add(user)
    db.commit()
    return create_access_token({"sub": str(user.id), "role": user.role, "jti": str(uuid.uuid4())})


@pytest.fixture
def syllabus(db, instructor_user):
    course = Course(title="DSA", code="ITC 212", academic_period="AY 2025-2026 S1",
                    instructor_id=instructor_user.id)
    db.add(course)
    db.flush()
    syl = Syllabus(course_id=course.id, file_path="x", status=SyllabusStatus.pending)
    db.add(syl)
    db.flush()
    db.add(Competency(syllabus_id=syl.id, competency_code="DSA-01", topic_name="Big-O",
                      learning_outcomes=["Analyze running time"]))
    db.commit()
    return syl


EDITED = {"competencies": [
    {"competency_code": "DSA-01", "topic_name": "Algorithm Analysis", "learning_outcomes": ["Analyze running time"]},
    {"competency_code": "DSA-02", "topic_name": "Linked Lists", "learning_outcomes": ["Implement a linked list"]},
]}


def _endorse(client, db, syllabus_id):
    for role, email in [(UserRole.program_director, "pd@test.com"), (UserRole.dean, "dean@test.com"),
                        (UserRole.vp_academic_affairs, "vp@test.com")]:
        token = _approver(db, role, email)
        r = client.post(f"/syllabi/{syllabus_id}/approve", headers=_auth(token), json={"feedback": "ok"})
        assert r.status_code == 200, r.text


def test_instructor_edits_draft_competencies(client, instructor_token, syllabus):
    r = client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(instructor_token), json=EDITED)
    assert r.status_code == 200, r.text
    assert sorted(c["competency_code"] for c in r.json()["competencies"]) == ["DSA-01", "DSA-02"]


def test_edit_rejects_duplicate_codes(client, instructor_token, syllabus):
    body = {"competencies": [EDITED["competencies"][0], {**EDITED["competencies"][1], "competency_code": "dsa-01"}]}
    r = client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(instructor_token), json=body)
    assert r.status_code == 422


def test_edit_blocked_for_other_instructor(client, other_instructor_token, syllabus):
    r = client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(other_instructor_token), json=EDITED)
    assert r.status_code == 403


def test_competencies_locked_once_submitted(client, instructor_token, syllabus):
    assert client.post(f"/syllabi/{syllabus.id}/submit-for-review", headers=_auth(instructor_token)).status_code == 200
    r = client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(instructor_token), json=EDITED)
    assert r.status_code == 409


def test_vp_endorsement_validates_and_locks(client, db, instructor_token, syllabus):
    client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(instructor_token), json=EDITED)
    client.post(f"/syllabi/{syllabus.id}/submit-for-review", headers=_auth(instructor_token))
    _endorse(client, db, syllabus.id)

    db.expire_all()
    syl = db.get(Syllabus, syllabus.id)
    assert syl.approval_stage == ApprovalStage.fully_approved
    assert syl.status == SyllabusStatus.validated
    comps = client.get(f"/syllabi/courses/{syl.course_id}/competencies", headers=_auth(instructor_token)).json()
    assert len(comps) == 2

    # No edits after endorsement, through either endpoint
    assert client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(instructor_token),
                      json=EDITED).status_code == 409
    changed = {"competencies": [{**EDITED["competencies"][0], "topic_name": "Something else"},
                                EDITED["competencies"][1]]}
    assert client.put(f"/syllabi/{syllabus.id}/validate", headers=_auth(instructor_token),
                      json=changed).status_code == 409
    # Confirming without changes is fine
    assert client.put(f"/syllabi/{syllabus.id}/validate", headers=_auth(instructor_token), json={}).status_code == 200


def test_rejected_syllabus_can_be_edited_again(client, db, instructor_token, syllabus):
    client.post(f"/syllabi/{syllabus.id}/submit-for-review", headers=_auth(instructor_token))
    pd = _approver(db, UserRole.program_director, "pd2@test.com")
    assert client.post(f"/syllabi/{syllabus.id}/reject", headers=_auth(pd),
                       json={"feedback": "Split topic 1"}).status_code == 200
    r = client.put(f"/syllabi/{syllabus.id}/competencies", headers=_auth(instructor_token), json=EDITED)
    assert r.status_code == 200


def _upload(client, token, course_id):
    import io
    from docx import Document
    doc = Document()
    doc.add_paragraph("Course Learning Outcomes")
    doc.add_paragraph("CLO1: Analyze the running time of algorithms using Big-O notation")
    doc.add_paragraph("CLO2: Implement linked lists, stacks and queues")
    buf = io.BytesIO()
    doc.save(buf)
    return client.post(f"/syllabi/upload?course_id={course_id}", headers=_auth(token),
                       files={"file": ("s.docx", buf.getvalue(),
                                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_reupload_replaces_draft_and_is_refused_under_review(client, db, instructor_token, syllabus, monkeypatch):
    from backend.core.limiter import limiter
    monkeypatch.setattr(limiter, "enabled", False)
    course_id = syllabus.course_id
    r = _upload(client, instructor_token, course_id)
    assert r.status_code == 201, r.text
    assert r.json()["id"] == str(syllabus.id)                       # replaced in place
    assert db.query(Syllabus).filter(Syllabus.course_id == course_id).count() == 1
    client.post(f"/syllabi/{syllabus.id}/submit-for-review", headers=_auth(instructor_token))
    assert _upload(client, instructor_token, course_id).status_code == 409


def test_duplicate_course_code_and_period_rejected(client, instructor_token, syllabus):
    body = {"title": "DSA again", "code": "itc212", "academic_period": "AY 2025-2026 S1"}
    r = client.post("/courses/", headers=_auth(instructor_token), json=body)
    assert r.status_code == 409


def test_second_copy_of_same_course_cannot_enter_review(client, db, instructor_token, instructor_user, syllabus):
    r = client.post(f"/syllabi/{syllabus.id}/submit-for-review", headers=_auth(instructor_token))
    assert r.status_code == 200, r.text
    # A pre-existing duplicate course row (created before the course-level guard) with its own draft syllabus
    twin_course = Course(title="DSA", code="ITC 212", academic_period="AY 2025-2026 S1",
                         instructor_id=instructor_user.id)
    db.add(twin_course)
    db.flush()
    twin = Syllabus(course_id=twin_course.id, file_path="y", status=SyllabusStatus.pending)
    db.add(twin)
    db.commit()
    r = client.post(f"/syllabi/{twin.id}/submit-for-review", headers=_auth(instructor_token))
    assert r.status_code == 409
    # Re-submitting the same syllabus (double click) is also refused
    r = client.post(f"/syllabi/{syllabus.id}/submit-for-review", headers=_auth(instructor_token))
    assert r.status_code == 400


def test_approver_can_open_original_file(client, db, syllabus):
    import os
    from backend.routers.syllabi import UPLOAD_DIR
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, f"{syllabus.course_id}_DSA.pdf")
    with open(path, "wb") as f:
        f.write(b"%PDF-1.4 test")
    try:
        syllabus.file_path = path
        db.commit()
        pd = _approver(db, UserRole.program_director, "pd.file@test.com")
        r = client.get(f"/syllabi/{syllabus.id}/file", headers=_auth(pd))
        assert r.status_code == 200 and r.content == b"%PDF-1.4 test"
        assert 'filename="DSA.pdf"' in r.headers["content-disposition"]

        other = User(email="pd.cba@test.com", password_hash=hash_password("Approver@1234"), full_name="CBA",
                     role=UserRole.program_director, is_active=True, college="CBA")
        db.add(other)
        db.commit()
        tok = create_access_token({"sub": str(other.id), "role": other.role, "jti": str(uuid.uuid4())})
        assert client.get(f"/syllabi/{syllabus.id}/file", headers=_auth(tok)).status_code == 403

        syllabus.file_path = "/etc/hosts"   # never serve files outside the upload folder
        db.commit()
        assert client.get(f"/syllabi/{syllabus.id}/file", headers=_auth(pd)).status_code == 404
    finally:
        os.remove(path)
