"""
Mapping validation (panel comment 1): stratified sample, blind labeling by two or three
validators, majority-vote ground truth with a Program Chair tie-break, HITL routing, and
accuracy / kappa / confidence-band metrics.
"""
import io
import uuid
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from backend.core.security import create_access_token, hash_password
from backend.models.assessment import (
    Assessment, AssessmentApprovalStatus, AssessmentItem, AssessmentType, GradingPeriod,
)
from backend.models.course import Course
from backend.models.syllabus import Competency, Syllabus, SyllabusStatus
from backend.models.user import User, UserRole
from backend.services.mapping_validation import _allocate, cohens_kappa, draw_sample, fleiss_kappa, wilson_interval


def _auth(user):
    return {"Authorization": "Bearer " + create_access_token(
        {"sub": str(user.id), "role": user.role, "jti": str(uuid.uuid4())})}


def _user(db, email, role, college="ICS"):
    u = User(email=email, password_hash=hash_password("Pass@1234"), full_name=email.split("@")[0].title(),
             role=role, is_active=True, college=college)
    db.add(u)
    db.commit()
    return u


# ── pure functions ──

def test_cohens_kappa_known_values():
    assert cohens_kappa([("a", "a"), ("b", "b")]) == 1.0
    # textbook example: 20 yes/yes, 5 yes/no, 10 no/yes, 15 no/no -> kappa 0.4
    pairs = [("y", "y")] * 20 + [("y", "n")] * 5 + [("n", "y")] * 10 + [("n", "n")] * 15
    assert cohens_kappa(pairs) == 0.4
    assert cohens_kappa([]) is None


def test_fleiss_kappa_known_values():
    assert fleiss_kappa([("a", "a", "a"), ("b", "b", "b")]) == 1.0
    # 2 of 3 agree on every item, labels balanced -> P̄ = 1/3, Pe = 1/2 -> kappa -1/3
    assert fleiss_kappa([("a", "a", "b"), ("b", "b", "a")]) == -0.333
    assert fleiss_kappa([]) is None


def test_wilson_interval():
    lo, hi = wilson_interval(85, 100)
    assert (lo, hi) == (76.7, 90.7)
    assert wilson_interval(0, 0) == (None, None)


def test_allocation_is_proportional_and_capped():
    assert _allocate({"a": 100, "b": 50}, 30) == {"a": 20, "b": 10}
    assert _allocate({"a": 5, "b": 3}, 30) == {"a": 5, "b": 3}
    assert sum(_allocate({"a": 7, "b": 7, "c": 7}, 10).values()) == 10


def test_sample_is_reproducible_and_stratified():
    class It:
        def __init__(self, i, conf):
            self.id, self.ai_confidence, self.confidence_score = f"i{i:03d}", conf, conf
    items = [It(i, 0.9 if i % 4 else 0.3) for i in range(80)]
    course_of = {it.id: ("c1" if i < 60 else "c2") for i, it in enumerate(items)}
    s1 = draw_sample(items, course_of, 20, seed=7, threshold=0.75)
    s2 = draw_sample(items, course_of, 20, seed=7, threshold=0.75)
    assert [x[0].id for x in s1] == [x[0].id for x in s2]
    assert sum(1 for _, c, _ in s1 if c == "c1") == 15 and sum(1 for _, c, _ in s1 if c == "c2") == 5
    assert sum(1 for _, _, b in s1 if b == "low") == 5   # 25% of items are low confidence


# ── end to end ──

@pytest.fixture
def setup(db, instructor_user):
    pd = _user(db, "pd@test.com", UserRole.program_director)
    second = _user(db, "second@test.com", UserRole.instructor)
    course = Course(title="DSA", code="ITC 212", academic_period="S1", instructor_id=instructor_user.id)
    db.add(course)
    db.flush()
    syl = Syllabus(course_id=course.id, file_path="x", status=SyllabusStatus.validated)
    db.add(syl)
    db.flush()
    comps = [Competency(syllabus_id=syl.id, competency_code=f"DSA-0{i}", topic_name=f"Topic {i}", learning_outcomes=[])
             for i in (1, 2)]
    db.add_all(comps)
    a = Assessment(course_id=course.id, title="Prelim", type=AssessmentType.exam, grading_period=GradingPeriod.prelim,
                   approval_status=AssessmentApprovalStatus.approved)
    db.add(a)
    db.flush()
    now = datetime.now(timezone.utc)
    # q0, q1: AI says DSA-01 (high conf); q2: AI says DSA-02 (low conf); q3: pre-tagged (not eligible)
    specs = [("q0", comps[0], 0.9), ("q1", comps[0], 0.95), ("q2", comps[1], 0.4)]
    for text, comp, conf in specs:
        db.add(AssessmentItem(assessment_id=a.id, question_text=text, mapped_competency_id=comp.id,
                              ai_suggested_competency_id=comp.id, ai_confidence=conf, confidence_score=conf,
                              ai_mapped_at=now))
    db.add(AssessmentItem(assessment_id=a.id, question_text="q3", mapped_competency_id=comps[0].id, confidence_score=1.0))
    db.commit()
    return {"pd": pd, "second": second, "course": course, "comps": comps}


def test_full_validation_flow(client, db, instructor_user, setup):
    pd, second, course, comps = setup["pd"], setup["second"], setup["course"], setup["comps"]

    # instructors cannot create sets
    assert client.post("/mapping-validation/", headers=_auth(instructor_user),
                       json={"name": "Pilot", "course_ids": [str(course.id)]}).status_code == 403

    r = client.post("/mapping-validation/", headers=_auth(pd),
                    json={"name": "Pilot validation", "course_ids": [str(course.id)], "sample_size": 10, "seed": 42})
    assert r.status_code == 201, r.text
    s = r.json()
    assert s["items"] == 3                                   # pre-tagged item excluded
    assert s["courses"][0]["validator_a"]["id"] == str(instructor_user.id)
    assert s["courses"][0]["adjudicator"]["id"] == str(pd.id)
    set_id, cid = s["id"], str(course.id)

    r = client.put(f"/mapping-validation/{set_id}/assignments/{cid}", headers=_auth(pd), json={
        "validator_a_id": str(instructor_user.id), "validator_b_id": str(second.id), "adjudicator_id": str(pd.id)})
    assert r.status_code == 200, r.text

    # blind queue: no AI fields anywhere
    q = client.get(f"/mapping-validation/{set_id}/label/{cid}", headers=_auth(instructor_user)).json()
    assert q["total"] == 3 and q["done"] == 0
    assert "ai" not in str(q).lower().replace("ainment", "").replace("again", "")  # no ai_* keys
    assert all(set(i) == {"id", "question_text", "my_label", "locked"} for i in q["items"])
    # the PD (not a validator) cannot open the labeling queue
    assert client.get(f"/mapping-validation/{set_id}/label/{cid}", headers=_auth(pd)).status_code == 403

    by_text = {i["question_text"]: i["id"] for i in q["items"]}
    c1, c2 = str(comps[0].id), str(comps[1].id)
    lab = lambda user, item, body: client.put(f"/mapping-validation/{set_id}/label/items/{by_text[item]}",
                                              headers=_auth(user), json=body)
    # q0: both say DSA-01 (AI right). q1: both say DSA-02 (AI wrong). q2: they disagree.
    for u in (instructor_user, second):
        assert lab(u, "q0", {"competency_id": c1}).status_code == 200
        assert lab(u, "q1", {"competency_id": c2}).status_code == 200
    lab(instructor_user, "q2", {"competency_id": c2})
    lab(second, "q2", {"none_fits": True})
    assert lab(second, "q2", {"competency_id": str(uuid.uuid4())}).status_code == 422

    res = client.get(f"/mapping-validation/{set_id}", headers=_auth(pd)).json()["metrics"]
    assert res["ground_truth_items"] == 2 and res["awaiting_adjudication"] == 1
    assert res["accuracy_pct"] == 50.0
    # validators never get results
    assert "metrics" not in client.get(f"/mapping-validation/{set_id}", headers=_auth(second)).json()

    adj = client.get(f"/mapping-validation/{set_id}/adjudication", headers=_auth(pd)).json()
    rows = adj["courses"][0]["items"]
    assert len(rows) == 1 and rows[0]["question_text"] == "q2"
    assert {l["label"] for l in rows[0]["labels"]} == {"DSA-02 · Topic 2", "No matching competency"}
    assert client.put(f"/mapping-validation/{set_id}/adjudication/items/{rows[0]['id']}", headers=_auth(pd),
                      json={"competency_id": c2, "note": "Matches outcome 2"}).status_code == 200
    # labels are final after adjudication
    assert lab(second, "q2", {"competency_id": c1}).status_code == 409

    m = client.get(f"/mapping-validation/{set_id}", headers=_auth(pd)).json()["metrics"]
    assert m["complete"] and m["ground_truth_items"] == 3 and m["correct"] == 2
    assert m["accuracy_pct"] == 66.7 and m["meets_target"] is False
    assert m["by_band"]["high"] == {**m["by_band"]["high"], "items": 2, "correct": 1}
    assert m["by_band"]["low"]["items"] == 1 and m["by_band"]["low"]["correct"] == 1
    assert m["kappa"] is not None and m["agreement_pct"] == 66.7

    x = client.get(f"/mapping-validation/{set_id}/export", headers=_auth(pd))
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Items", "By Competency"]
    assert wb["Items"].max_row == 4


def test_validator_cannot_be_replaced_after_labeling(client, db, instructor_user, setup):
    pd, second, course = setup["pd"], setup["second"], setup["course"]
    s = client.post("/mapping-validation/", headers=_auth(pd),
                    json={"name": "Pilot", "course_ids": [str(course.id)], "sample_size": 10, "seed": 1}).json()
    q = client.get(f"/mapping-validation/{s['id']}/label/{course.id}", headers=_auth(instructor_user)).json()
    client.put(f"/mapping-validation/{s['id']}/label/items/{q['items'][0]['id']}", headers=_auth(instructor_user),
               json={"none_fits": True})
    r = client.put(f"/mapping-validation/{s['id']}/assignments/{course.id}", headers=_auth(pd),
                   json={"validator_a_id": str(second.id), "validator_b_id": None, "adjudicator_id": str(pd.id)})
    assert r.status_code == 409


def test_pd_outside_college_cannot_create(client, db, setup):
    other_pd = _user(db, "pd2@test.com", UserRole.program_director, college="CBA")
    r = client.post("/mapping-validation/", headers=_auth(other_pd),
                    json={"name": "Other college", "course_ids": [str(setup["course"].id)]})
    assert r.status_code == 403


def test_report_includes_validation_results(client, db, instructor_user, admin_user, setup):
    from backend.services.reports.data import ReportFilters, build_report
    pd, course = setup["pd"], setup["course"]
    client.post("/mapping-validation/", headers=_auth(pd),
                json={"name": "Pilot set", "course_ids": [str(course.id)], "sample_size": 10, "seed": 3})
    data = build_report(db, admin_user, ReportFilters())
    assert [v["name"] for v in data["mapping_validation"]] == ["Pilot set"]
    assert data["mapping_validation"][0]["items"] == 3
    pdf = client.get("/reports/academic-performance/export", params={"format": "pdf"}, headers=_auth(admin_user))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_three_validators_majority_and_hitl_routing(client, db, instructor_user, setup):
    pd, second, course, comps = setup["pd"], setup["second"], setup["course"], setup["comps"]
    third = _user(db, "third@test.com", UserRole.instructor)
    s = client.post("/mapping-validation/", headers=_auth(pd),
                    json={"name": "Capstone protocol", "course_ids": [str(course.id)], "seed": 3}).json()
    set_id, cid = s["id"], str(course.id)
    # validators must be three different people, and the Program Chair is none of them
    body = {"validator_a_id": str(instructor_user.id), "validator_b_id": str(second.id),
            "validator_c_id": str(second.id), "adjudicator_id": str(pd.id)}
    assert client.put(f"/mapping-validation/{set_id}/assignments/{cid}", headers=_auth(pd), json=body).status_code == 422
    body["validator_c_id"] = str(third.id)
    r = client.put(f"/mapping-validation/{set_id}/assignments/{cid}", headers=_auth(pd), json=body)
    assert r.status_code == 200 and r.json()["courses"][0]["validator_c"]["id"] == str(third.id)

    q = client.get(f"/mapping-validation/{set_id}/label/{cid}", headers=_auth(third)).json()
    by_text = {i["question_text"]: i["id"] for i in q["items"]}
    c1, c2 = str(comps[0].id), str(comps[1].id)
    lab = lambda user, item, body: client.put(f"/mapping-validation/{set_id}/label/items/{by_text[item]}",
                                              headers=_auth(user), json=body)
    live = lambda text: db.query(AssessmentItem).filter(AssessmentItem.question_text == text).one()
    for text in ("q0", "q1"):  # faculty already confirmed the AI's mapping
        it = live(text)
        it.needs_review, it.validated_by = False, instructor_user.id
    db.commit()

    # q0: unanimous DSA-01 (AI right). q1: 2 of 3 say DSA-02 (AI wrong). q2: three different answers.
    for u in (instructor_user, second, third):
        lab(u, "q0", {"competency_id": c1})
    lab(instructor_user, "q1", {"competency_id": c2})
    lab(second, "q1", {"competency_id": c1})
    lab(instructor_user, "q2", {"competency_id": c1})
    lab(second, "q2", {"competency_id": c2})
    # A and B differ, C hasn't answered: not a tie yet
    assert client.get(f"/mapping-validation/{set_id}/adjudication", headers=_auth(pd)).json()["courses"][0]["items"] == []
    lab(third, "q2", {"none_fits": True})
    lab(third, "q1", {"competency_id": c2})

    db.expire_all()
    assert live("q0").needs_review is False                          # AI matched ground truth
    assert live("q1").needs_review is True and live("q1").validated_by is None   # AI wrong -> HITL

    rows = client.get(f"/mapping-validation/{set_id}/adjudication", headers=_auth(pd)).json()["courses"][0]["items"]
    assert len(rows) == 1 and rows[0]["question_text"] == "q2" and len(rows[0]["labels"]) == 3
    assert client.put(f"/mapping-validation/{set_id}/adjudication/items/{rows[0]['id']}", headers=_auth(pd),
                      json={"competency_id": c1}).status_code == 200

    m = client.get(f"/mapping-validation/{set_id}", headers=_auth(pd)).json()["metrics"]
    assert m["complete"] and m["correct"] == 1 and m["accuracy_pct"] == 33.3
    assert m["decisions"] == {"unanimous": 1, "majority": 1, "program_chair": 1}
    assert m["kappa_method"] == "Fleiss" and m["labeled_c"] == 3
    db.expire_all()
    assert live("q2").needs_review is True                            # chair says DSA-01, AI said DSA-02

    x = client.get(f"/mapping-validation/{set_id}/export", headers=_auth(pd))
    ws = load_workbook(io.BytesIO(x.content))["Items"]
    assert [c.value for c in ws[1]][5:11] == ["Validator A", "Validator B", "Validator C", "Program Chair",
                                              "Ground truth", "Decision"]
